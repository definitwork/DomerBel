import time
import openpyxl
from advertisement.models import Advertisement, Category, Region, Spisok, ElementTwo, PhotoAdvertisement, ErrorFile
from advertisement.validators import validate_words
from users.models import User
from transliterate import slugify
from zipfile import ZipFile
from advertisement.utils_for_models import upload_to
import os
import xlsxwriter
import shutil


def check_article(ads,value_author):
    advertisements = Advertisement.objects.filter(author=value_author)
    '''Функция, проверяющаяя артикул в объявлениях'''
    if 'артикул' in ads:
        article = ads.get('артикул')
        if article == None:
            return True
        elif article != None and len(str(article)) < 256:
            articals = []
            for i in advertisements:
                articals.append(i.article)
            if article in articals:
                return f'Уже существует объявление с таким артикулом {article}'
            else:
                return True
        elif article != None and len(str(article)) > 255:
            return 'Количество символов в артикле не должно быть больше 255'
    else:
        return True

def chek_title(ads):
    '''Функция, проверяющая заголовок объявления'''
    if 'заголовок' in ads and len(str(ads.get("заголовок"))) < 256:
        try:
            validate_words(str(ads.get("заголовок")))
            return True
        except:
            return 'Найдено запрещенное слово'
    else:
        return 'поле "Заголовк" обязательное'

def chek_price(ads):
    '''Функция, проверящя цену'''
    if 'цена' in ads and ads.get('цена') != None:
        price = str(ads.get('цена'))
        if price.find('.') != -1:
            list_price = price.split('.')
            if len(list_price) == 2 and len(list_price[0])+len(list_price[1]) <= 12 and len(list_price[1]) <= 2 and list_price[0].isdigit() and list_price[1].isdigit():
                return True
            else:
                return 'некорретно указана цена'
        else:
            if price.isdigit() and len(price) <= 12:
                return True
            else:
                return 'некорретно указана цена'
    elif ads.get('цена') == None:
        return True
    else:
        return True

def chek_category(ads):
    '''Функция, проверяющая категорию объявления'''
    if 'категория' in ads:
        category = ads.get('категория')
        if Category.objects.filter(title__iexact = f'{category}'):
            return True
        else:
            return 'некорректно заполнено поле "Категория"'
    else:
        return 'поле "Категория" обязательное'

def check_region(ads):
    '''Функция, проверяющаа регион'''
    if 'регион' in ads:
        region = ads.get('регион')
        if Region.objects.filter(area__iexact=f'{region}'):
            return True
        else:
            return 'некорректно заполнено поле "Регион"'
    else:
        return 'поле "Регион" обязательное'

def check_additional_information(ads,check_category):
    '''Функция проверяет поля для детальной информации и возвращает JsonFile'''
    data = {}
    error = []
    if check_category == True:
        if 'категория' in ads:
            category_from_ads = ads.get('категория')
            category = Category.objects.prefetch_related('field_set').get(title=f'{category_from_ads }')
            '''Все поля'''
            fields = category.field_set.all()
            for field_from_excel in ads.keys():
                if field_from_excel not in ['артикул','заголовок','категория', 'регион','описание'] and not field_from_excel.startswith('фото'):
                    value = ads.get(field_from_excel)
                    field_from_db = fields.filter(title__iexact = field_from_excel)
                    if len(field_from_db) != 0:
                        for field in field_from_db:
                            if field.error != '':
                                '''обязательные поля для заполнения'''
                                if field.spisok_id != None:
                                    '''обязательные поля с варинтами для заполнения'''
                                    spisok = Spisok.objects.prefetch_related('element_set').get(id=field.spisok_id)
                                    elements_from_db = spisok.element_set.all()
                                    elements_for_compare = elements_from_db.filter(title__iexact = ads.get(field_from_excel))
                                    if len(elements_for_compare) != 0:
                                        '''если значение поля состоит из одной части'''
                                        for element in elements_for_compare:
                                            data[field.title] = element.title
                                    else:
                                        '''если значение поля стостоит из двух частей или некорректно введено'''
                                        value_from_ads = ads.get(field.title.lower())
                                        item_for_find = value_from_ads.split(' ',1)[0]
                                        element_one = elements_from_db.filter(title__istartswith = item_for_find)
                                        id_elem_one= [item.id for item in element_one]
                                        element_two = ElementTwo.objects.filter(element_id__in = id_elem_one)
                                        if len(element_two) != 0:
                                            '''В словаре все возможные сочетания элементов Element и ElementTwo'''
                                            option = {}
                                            for one in element_one:
                                                for two in element_two:
                                                    key = one.title.lower()+' '+two.title.lower()
                                                    value = [one,two]
                                                    option[key]=value
                                            element_one_two = option.get(value_from_ads.lower())
                                            if element_one_two != None:
                                                data[field.title] = element_one_two[0].title+ ', ' + element_one_two[1].title
                                            else:
                                                error.append(f'Некорректное значение "{value_from_ads}"')
                                        else:
                                            error.append(f'Некорректное значение "{value_from_ads}"')
                                else:
                                    ''' обязательные поля без вариантов для заполнения'''
                                    data[field.title] = value
                            else:
                                '''необязательные полян для заполнения'''
                                data[field.title] = value
                    else:
                        error.append(f'Некорректное имя поля "{field_from_excel}"')
        else:
            return ['не заполнено поле категория']
        if len(error) != 0:
            return error
        else:
            return data
    else:
        return ['некорректно заполнено поле категории']

def chek_description(ads):
    '''Функция, проверяющаа описание'''
    if 'описание' in ads:
        try:
            validate_words(str(ads.get("описание")))
            return True
        except:
            return 'Найдено запрещенное слово'
    else:
        return 'поле "Описание" обязательное'

def update_photo(ads_for_save,file_name,uploud_zip,email):
    '''Функция для извлечения preview_image из электронного архива,
    и хэширование имени файла и распределение файлов по приложениям
    и далее в разные папки случайным образом'''
    try:
        with ZipFile(uploud_zip,'r') as zip:
            image_from_zip = zip.extract(file_name,f'./media/files_for_bulk_import_of_ads/{email}')
        new_location_image = upload_to(ads_for_save,image_from_zip)
        preview_image = os.replace(f'./{image_from_zip}',f'./media/{new_location_image}')
        return f'{new_location_image}'
    except:
        return False

def write_file_with_error_ads(list_error,email, user):
    path = f'./media/files_for_bulk_import_of_ads/{email}/error_{email}_{time.time()}.xlsx'
    book = xlsxwriter.Workbook(path)
    sheet = book.add_worksheet()
    field = {}
    for row in range(0, len(list_error) + 1):
        if row == 0:
            values = list_error[0]
            colum = 0
            for title in values.keys():
                sheet.write(row, colum, title)
                field[title] = colum
                colum = colum + 1
        elif row == 1:
            values = list_error[0]
            for title in values.keys():
                colum = field.get(title)
                sheet.write(row, colum, values.get(title))
        else:
            values = list_error[row - 1]
            for title in values.keys():
                if title in field.keys():
                    colum = field.get(title)
                    sheet.write(row, colum, values.get(title))
                else:
                    field[title] = len(field)
                    colum = field.get(title)
                    sheet.write(0, colum, title)
                    sheet.write(row, colum, values.get(title))
    book.close()
    file_error = ErrorFile(file = path, user = user)
    file_error.save()
    return file_error

def save_many_ads_from_excel(uploud_file,id,first_name,phone_number,email):
    '''Функция сохраняющаяя обявления из экселя'''
    book = openpyxl.open(uploud_file, read_only=True)
    sheet = book.active
    row_in_excel = 2
    list_ads_error = []
    value_author = User.objects.get(id=id)
    while row_in_excel <= sheet._max_row:
        ads = {}
        for i in range(0, sheet.max_column):
            if sheet[row_in_excel][i].value != None:
                ads[f'{sheet[1][i].value.lower()}'] = sheet[row_in_excel][i].value
        if len(ads) != 0:
            status_ads = []
            value_artical = check_article(ads,value_author)
            if value_artical != True:
                status_ads.append(value_artical)
            value_title = chek_title(ads)
            if value_title != True:
                status_ads.append(value_title)
            value_price = chek_price(ads)
            if value_price != True:
                status_ads.append(value_price)
            value_category = chek_category(ads)
            if value_category != True:
                status_ads.append(value_category)
            value_region = check_region(ads)
            if value_region != True:
                status_ads.append(value_region)
            json_file = check_additional_information(ads,value_category)
            if type(json_file) == list:
                for item in json_file:
                    status_ads.append(item)
            value_description = chek_description(ads)
            if value_description != True:
                status_ads.append(value_description)
            if len(status_ads) == 0:
                ads_for_save = Advertisement(author = value_author,
                article = str(ads.get('артикул')),
                title = str(ads.get('заголовок')),
                price = float(ads.get('цена')),
                category = Category.objects.get(title__iexact=ads.get('категория')),
                bearer = 'Компания',
                region = Region.objects.get(area__iexact=ads.get('регион')),
                contact_name = first_name,
                phone_num = phone_number,
                email = email,
                slug = slugify(f"{ads.get('заголовок')}__{ads.get('артикул')}"),
                additional_information = json_file,
                description = str(ads.get('описание')))
                try:
                    ads_for_save.save()
                except:
                    list_ads_error.append(ads)
            else:
                list_ads_error.append(ads)
        row_in_excel = row_in_excel + 1
    if len(list_ads_error) != 0:
        file_error = write_file_with_error_ads(list_ads_error,email,value_author)
        return {'file': [file_error.id,file_error.file]}
    else:
        return True

def save_many_ads_from_zip(uploud_zip,id,first_name,phone_number,email):
    '''Функция для сохранениея обявлений из электронного архива'''
    list_ads_error = []
    value_author = User.objects.get(id=id)
    with ZipFile(uploud_zip,'r') as zip:
        for file in zip.namelist():
            if file.endswith('.xlsx'):
                excel = zip.extract(file,path=f'./media/files_for_bulk_import_of_ads/{email}/')
    book = openpyxl.open(excel, read_only=True)
    sheet = book.active
    row_in_excel = 2
    while row_in_excel <= sheet._max_row:
        ads = {}
        for i in range(0, sheet.max_column):
            if sheet[row_in_excel][i].value != None:
                ads[f'{sheet[1][i].value.lower()}'] = sheet[row_in_excel][i].value
        if len(ads) != 0:
            status_ads = []
            value_artical = check_article(ads,value_author)
            if value_artical != True:
                status_ads.append(value_artical)
            value_title = chek_title(ads)
            if value_title != True:
                status_ads.append(value_title)
            value_price = chek_price(ads)
            if value_price != True:
                status_ads.append(value_price)
            value_category = chek_category(ads)
            if value_category != True:
                status_ads.append(value_category)
            value_region = check_region(ads)
            if value_region != True:
                status_ads.append(value_region)
            json_file = check_additional_information(ads,value_category)
            if type(json_file) == list:
                for item in json_file:
                    status_ads.append(item)
            value_description = chek_description(ads)
            if value_description != True:
                status_ads.append(value_description)
            if len(status_ads) == 0:
                ads_for_save = Advertisement(author = value_author,
                article = str(ads.get('артикул')),
                title = str(ads.get('заголовок')),
                price = float(ads.get('цена')),
                category = Category.objects.get(title__iexact=ads.get('категория')),
                bearer = 'Компания',
                region = Region.objects.get(area__iexact=ads.get('регион')),
                contact_name = first_name,
                phone_num = phone_number,
                email = email,
                additional_information = json_file,
                description = str(ads.get('описание')))
                ads_for_save.save()
                preview_image = update_photo(ads_for_save, ads.get('фото1'), uploud_zip, email)
                if preview_image != False:
                    ads_for_save.preview_image = preview_image
                    ads_for_save.save(update_fields=["preview_image"])
                    photo_in_db=[]
                    photo_save = PhotoAdvertisement(
                                    photo = preview_image,
                                    advertisement = ads_for_save)
                    photo_save.save()
                    photo_in_db.append(photo_save)
                    for key in ads.keys():
                        if key.startswith('фото') and key != 'фото1':
                            photo = update_photo(ads_for_save, ads.get(key), uploud_zip, email)
                            if photo != False:
                                photo_save = PhotoAdvertisement(
                                    photo=photo,
                                    advertisement=ads_for_save)
                                photo_save.save()
                                photo_in_db.append(photo_save)

                            else:
                                for photo in photo_in_db:
                                    os.remove(f'./media/{photo.photo}')
                                ads_for_save.delete()
                                status_ads.append(f'Невозможно сохранить изображение {ads.get(key)}')
                                list_ads_error.append(ads)
                                break
                else:
                    ads_for_save.delete()
                    status_ads.append(f'Невозможно сохранить файл {ads.get("фото1")}')
                    list_ads_error.append(ads)
            else:
                list_ads_error.append(ads)
        row_in_excel = row_in_excel + 1
    if len(list_ads_error) != 0:
        file_error = write_file_with_error_ads(list_ads_error,email,value_author)
        return {'file': [file_error.id,file_error.file]}
    else:
        return True

def delete_everything_in_folder():
    '''Функция, удаляющая все файлы из папки для "files_for_bulk_import_of_ads" '''
    path = './media/files_for_bulk_import_of_ads'
    shutil.rmtree(path)
    os.mkdir(path)

